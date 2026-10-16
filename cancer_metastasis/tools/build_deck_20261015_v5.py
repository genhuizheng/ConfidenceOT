"""ConfidenceOT_20261015_v5.pptx: v4 plus the raw label's all-mice retained-fraction heatmap.

v4 is only read. v4 showed every mouse only for ranknm256_noscale_ds_cos (slides 23-24); the raw
label appeared only on the matched slides, restricted to the 8 mice with a ranking question.
v5 inserts the same two-part heatmap for raw, same layout, mice, candidate order, colour scale,
outlines and GT-rank rule, right after the main one:

23-24 Retained Fractions and Lineage Ground Truth (1/2, 2/2)          ranknm256_noscale_ds_cos (v4)
25-26 Retained Fractions and Lineage Ground Truth: raw (1/2, 2/2)     new
27-   the rest of v4, shifted by two; slide references in the notes updated.

    python cancer_metastasis/tools/build_deck_20261015_v5.py
"""

from __future__ import annotations

import importlib.util
import subprocess
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
_spec = importlib.util.spec_from_file_location("deck_v4", HERE / "build_deck_20261015_v4.py")
v4 = importlib.util.module_from_spec(_spec)
sys.modules["deck_v4"] = v4
_spec.loader.exec_module(v4)
v3, v2 = v4.v3, v4.v2
text, notes = v2.text, v2.notes
ORANGE, MUTED = v2.ORANGE, v2.MUTED
MAIN, RAW = v4.MAIN, v4.RAW

SRC = v4.OUT
OUT = Path(r"D:\Xia_lab\Progress_report\progress_report7\ConfidenceOT_20261015_v5.pptx")
TMP = OUT.with_name("_tmp_v5.pptx")
PARTS = ([["3451", "3454", "3457"], ["3460", "3465", "3466"]], [["3467", "3508", "3515"], ["3513", "3724"]])


def slide_heatmap_raw(slide, data: dict, rows_of_mice, part: str, hits: tuple, hits_main: tuple) -> None:
    v4.draw_rows(slide, data, rows_of_mice, 0.45, 0.95, 0.34, 0.5, 0.95)
    v4.legend(slide, 10.2, 1.0)
    text(slide, 10.2, 4.75, 2.8, 1.6, [[(f"{hits[0]} of {hits[1]}", {"bold": True, "size": 22, "color": v4.DARK_GREY})],
                                       [(f"lesions where a same-family primary ranks first under raw (all mice, ties "
                                         f"not counted); {hits_main[0]} of {hits_main[1]} under "
                                         f"{MAIN}", {"color": MUTED})]], size=12)
    text(slide, 0.45, 6.75, 12.4, 0.3, f"raw, M4-E; {part}. Source: readout/pairs_{RAW}.csv; same mice, lesions, "
         "candidate order, colour scale and outlines as slides 23–24.", size=10, color=MUTED, italic=True)
    notes(slide, "\n".join([
        "Real data, raw label. The same layout as slides 23–24 (ranknm256_noscale_ds_cos), so the two can be compared "
        "block by block: each block is one mouse; rows are its metastatic lesions, columns its candidate primaries in "
        "tumour order. Colour: retained fraction r on the same 0–1 scale. Black outline: same family (lineage ground "
        "truth). Hatched: not run (fewer than 20 malignant cells on a side).",
        "raw retains nearly every source cell in nearly every pair (median r 0.99), so almost every cell is at the dark "
        "end of the scale: the gate rejects little, and the ranking is decided by small differences in r. The "
        "uniform colour is the result, not missing data.",
        f"GT rank as on slides 23–24. Across both raw slides a same-family primary ranks first in {hits[0]} of "
        f"{hits[1]} lesions (strict: exact ties not counted), against {hits_main[0]} of {hits_main[1]} for "
        f"{MAIN}.",
        "The two labels differ in several preprocessing settings at once (slide 28). Exploratory reading; the "
        "prespecified readout is d_M.",
    ]))


def widen_title(slide) -> None:
    box = next(s for s in slide.shapes if s.name == "TextBox 3")
    right_edge = box.left + box.width
    box.left = Inches(3.65)
    box.width = right_edge - box.left
    for run in box.text_frame.paragraphs[0].runs:
        run.font.size = Pt(26)


def replace_in_notes(slide, old: str, new: str) -> None:
    runs = [r for p in slide.notes_slide.notes_text_frame.paragraphs for r in p.runs if old in r.text]
    assert len(runs) == 1, (old, len(runs))
    runs[0].text = runs[0].text.replace(old, new)


def main():
    data_main, data_raw = v3.heatmap_data(MAIN), v3.heatmap_data(RAW)
    hits_main, hits_raw = v4.tally(data_main), v4.tally(data_raw)

    script = v2.SKILL / "scripts" / "add_slide.py"
    subprocess.run([sys.executable, str(script), str(SRC), "slide13.xml", "-o", str(TMP)],
                   check=True, capture_output=True, text=True, encoding="utf-8")
    subprocess.run([sys.executable, str(script), str(TMP), "slide13.xml"],
                   check=True, capture_output=True, text=True, encoding="utf-8")
    prs = Presentation(str(TMP))
    slides = list(prs.slides)
    old, new = slides[:31], slides[31:]
    titles = [next(s for s in sl.shapes if s.name == "TextBox 3").text_frame.text for sl in old]
    assert titles[21].startswith("Tumor-Level Benchmark") and titles[23] == "Retained Fractions and Lineage Ground Truth (2/2)"
    assert titles[24].startswith("Does ConfidenceOT Rank") and titles[25].startswith("RNA Preprocessing")
    assert titles[26].startswith("Raw vs Preprocessed") and titles[27].startswith("Raw vs Preprocessed")

    for k, (slide, rows_of_mice) in enumerate(zip(new, PARTS), 1):
        v2.clear_content(slide)
        v2.set_title(slide, f"Retained Fractions and Lineage Ground Truth: raw ({k}/2)")
        widen_title(slide)
        slide_heatmap_raw(slide, data_raw, rows_of_mice, f"part {k} of 2", hits_raw, hits_main)

    # Slide references in the notes of slides that now sit two later.
    replace_in_notes(old[21], "slide 25.", "slide 27.")
    for slide in (old[26], old[27]):
        for paragraph in slide.notes_slide.notes_text_frame.paragraphs:
            for run in paragraph.runs:
                run.text = run.text.replace("(slide 26)", "(slide 28)")

    id_list = prs.slides._sldIdLst
    entries = list(id_list)
    by_slide = dict(zip([id(s) for s in slides], entries))
    order = old[:24] + new + old[24:]
    for entry in entries:
        id_list.remove(entry)
    for slide in order:
        id_list.append(by_slide[id(slide)])
    for number, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if shape.name == "Rectangle 6":
                for field in shape._element.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}fld"):
                    field.find("{http://schemas.openxmlformats.org/drawingml/2006/main}t").text = str(number)
    prs.save(str(OUT))
    TMP.unlink()
    print(OUT, len(prs.slides), "slides; strict hits main", hits_main, "raw", hits_raw)


if __name__ == "__main__":
    main()
