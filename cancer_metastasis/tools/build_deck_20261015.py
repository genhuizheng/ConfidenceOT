"""Build ConfidenceOT_20261015.pptx from ConfidenceOT_report_20261002.pptx.

Figures replaced by their bare versions (no titles, subtitles or reading notes),
GSEA as one signed dot plot per analysis, every explanatory text box moved into
the slide's notes. Tables stay. Old slides 8 and 9 are dropped (see NOTE_*).
"""
import copy
import re
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt

SRC = Path(r"D:\Xia_lab\Progress_report\progress_report7\ConfidenceOT_report_20261002.pptx")
OUT = Path(r"D:\Xia_lab\Progress_report\progress_report7\ConfidenceOT_20261015.pptx")
REPO = Path(r"D:\Xia_lab\CellOT")
FIG = REPO / "figures" / "slides_20261015"
DOT = REPO / "results" / "downstream_20261007" / "figures"

# Content area under the header and above the footer band.
AREA = (0.35, 0.78, 12.63, 6.27)

# old slide number -> new slide number; 8 and 9 are dropped.
NEW = {o: o for o in range(1, 8)}
NEW.update({o: o - 2 for o in range(10, 23)})

DOTPLOT_READING = (
    "Reading the figure: one signed GSEA per arm, primary rejected vs primary retained, retained as the "
    "reference. Colour: NES; blue = higher in retained cells, orange = higher in rejected cells; limits "
    "+/-2, points beyond take the end colour (exact NES in the CSV). Size: -log10 FDR in four bins "
    "(< 1, 1-2, 2-3, >= 3); FDR 0 (below the 2,000-permutation resolution) is drawn at 1/2000. Black "
    "outline: FDR < 0.05. Columns: all 50 hallmark sets, grouped by the eight hallmark process categories "
    "(Liberzon et al., Cell Systems 2015).")
# Written in the old numbering, like every text taken from the old deck; remap() converts it.
FACTORIAL_ROWS = "Rows: the 32 preprocessing arms of slide 5, at within-side acceptance 0.90 (the rerun of slide 13)."
CASES = ("Correct gate per case: Case 1, all populations shared: reject nothing. Case 2, one population "
         "unmatched: reject only that population. Case 3, no population shared: reject everything.")
UNIT_NES = ("In a few arms GSEApy returned NES of exactly +/-1.00 with p = 1 for many sets, never called: {}. "
            "They show as small pale dots.")
# Per new slide: (old text, new text) applied to the finished notes.
FIXES = {
    1: [("(a)", "(top row)"), ("(b)", "(bottom row)"), ("Left, ideal", "Left column, ideal"), ("Right, real", "Right column, real"),
        ("On the right the cells", "In the right column the cells")],
    2: [("Panel (a) is an illustrative draw.", "The upper part is an illustrative draw; columns: reference source, then the "
         "target at L0 matched, L1 depth mismatch, L2 depth + dropout.")],
    3: [("(a) population lost", "Top row, population lost"), ("(b) population emerged", "Bottom row, population emerged"),
        ("case (a)", "top row"), ("case (b)", "bottom row"), ("from the target (a) or from the source (b)",
         "from the target (top) or from the source (bottom)")],
    7: [("(next slide)", "(slides 11–14)")],
    9: [("Figure: method_comparison_heatmap.png (benchmark/55_plot_method_comparison.py), title block cropped. ",
         "Figure: benchmark/55_plot_method_comparison.py --bare. ")],
    11: [("What changed from slides 7–8", "What changed from the 10/02 run (slide 7)")],
    14: [("Framed: ", "")],
    16: [("What the two figures show", "What the figure shows")],
    17: [("What the two figures show", "What the figure shows")],
    18: [("What the two figures show", "What the figure shows")],
    19: [("(Table 5)", "(slide 15)")],
}
CYCLING_ROWS = ("Rows: RNA, RNA+CNV and CNV representations; ds_cos = ranknm256_noscale_ds_cos. Arms: A nothing "
                "removed; B malignant cells outside G1 removed on both sides; C the 97 cell-cycle genes removed; "
                "D both. CNV ran arms A and B only. Within-side acceptance 0.99.")


def remap(text: str) -> str:
    text = text.replace("slide 8 showed only the rows called in at least half of the arms, which are the top rows here",
                        "the 10/02 report showed only the rows called in at least half of the arms")
    text = text.replace("every retained-side pathway of slide 8 stays", "every retained-side pathway of the 10/02 run stays")
    text = text.replace("Reading slides 14–16 and 18–20", "Reading the GSEA slides")

    def one(match):
        word, a, b = match.group(1), int(match.group(2)), match.group(3)
        if b:
            return f"{word} {NEW.get(a, a)}–{NEW.get(int(b), int(b))}"
        return f"{word} {NEW.get(a, a)}"
    return re.sub(r"\b([Ss]lides?) (\d+)(?:[–-](\d+))?", one, text)


def text_of(shape) -> str:
    lines = []
    for paragraph in shape.text_frame.paragraphs:
        line = "".join(run.text for run in paragraph.runs).strip()
        if line:
            lines.append(line)
    return "\n".join(lines)


def shapes_named(slide, *names):
    return [shape for shape in slide.shapes if shape.name in names]


def remove(slide, shape):
    element = shape._element
    if shape.shape_type == 13:
        rid = element.blip_rId
        element.getparent().remove(element)
        if not any(getattr(s, "shape_type", None) == 13 and s._element.blip_rId == rid for s in slide.shapes):
            slide.part.drop_rel(rid)
    else:
        element.getparent().remove(element)


def take_text(slide, *names) -> list[str]:
    """Remove the named text shapes and return their text, in slide order."""
    texts = []
    for shape in list(slide.shapes):
        if shape.name in names and shape.has_text_frame:
            texts.append(text_of(shape))
            remove(slide, shape)
    return texts


def drop_pictures(slide, keep_logo=True):
    for shape in list(slide.shapes):
        if shape.shape_type == 13 and not (keep_logo and shape.name == "Picture 8" and shape.top < Inches(0.2)):
            remove(slide, shape)


def drop_highlights(slide):
    for shape in list(slide.shapes):
        if shape.name.startswith("Highlight "):
            remove(slide, shape)


def place(slide, path: Path, area=AREA, top_align=False):
    x0, y0, w0, h0 = area
    with Image.open(path) as image:
        ratio = image.width / image.height
    w, h = w0, w0 / ratio
    if h > h0:
        h, w = h0, h0 * ratio
    x = x0 + (w0 - w) / 2
    y = y0 if top_align else y0 + (h0 - h) / 2
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))


def move(shape, top=None, left=None):
    if top is not None:
        shape.top = Inches(top)
    if left is not None:
        shape.left = Inches(left)


def set_notes(slide, *parts):
    existing = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
    body = "\n\n".join(remap(p) for p in parts if p and p.strip())
    slide.notes_slide.notes_text_frame.text = (existing + "\n\n" + body).strip() if existing else body


def set_bullets(shape, heading: str, items: list[str]):
    """Keep the shape's first (heading) and second (bullet) paragraph formats; replace the text."""
    body = shape.text_frame._txBody
    paragraphs = body.findall("{http://schemas.openxmlformats.org/drawingml/2006/main}p")
    head, template = paragraphs[0], paragraphs[1]
    for p in paragraphs[1:]:
        body.remove(p)
    runs = head.findall("{http://schemas.openxmlformats.org/drawingml/2006/main}r")
    for extra in runs[1:]:
        head.remove(extra)
    runs[0].find("{http://schemas.openxmlformats.org/drawingml/2006/main}t").text = heading
    for item in items:
        p = copy.deepcopy(template)
        rs = p.findall("{http://schemas.openxmlformats.org/drawingml/2006/main}r")
        for extra in rs[1:]:
            p.remove(extra)
        rpr = rs[0].find("{http://schemas.openxmlformats.org/drawingml/2006/main}rPr")
        if rpr is not None:  # plain body font, whatever the first run was
            for attribute in ("b",):
                rpr.attrib.pop(attribute, None)
            for child in list(rpr):
                if child.tag.endswith("}latin") or child.tag.endswith("}cs") or child.tag.endswith("}ea"):
                    rpr.remove(child)
        rs[0].find("{http://schemas.openxmlformats.org/drawingml/2006/main}t").text = item
        body.append(p)


def main():
    prs = Presentation(str(SRC))
    slides = list(prs.slides)
    old = {i + 1: s for i, s in enumerate(slides)}

    # 1-3: figures
    for number, picture_name, text_name, figure in (
            (1, "Figure problem", "Text problem", FIG / "nuisance" / "problem.png"),
            (2, "Figure benchmark", "Text benchmark", FIG / "nuisance" / "benchmark.png"),
            (3, "Figure scenario", "Text scenario", FIG / "scenario" / "scenario_umap.png")):
        s = old[number]
        texts = take_text(s, text_name)
        for shape in shapes_named(s, picture_name):
            remove(s, shape)
        place(s, figure)
        set_notes(s, *texts, CASES if number in (2,) else "")

    # 4: table 1
    s = old[4]
    take_text(s, "TextBox 8")
    move(shapes_named(s, "Table setting")[0], top=0.95)
    set_notes(s, "Table 1. The technical benchmark: simulator, populations, sizes, replicates, the three technical "
                 "levels, the cases, preprocessing, the OT gate, the rejection cost and the run count.")

    # 5: preprocessing
    s = old[5]
    texts = take_text(s, "Text pipeline note", "Text arm names")
    move(shapes_named(s, "Table normalisation")[0], top=2.05)
    move(shapes_named(s, "Table options")[0], top=4.15)
    set_notes(s, *texts)

    # 6: arm heatmap
    s = old[6]
    texts = take_text(s, "Text heatmap")
    drop_pictures(s)
    place(s, FIG / "arm_heatmap" / "arm_error_heatmap.png")
    set_notes(s, *texts, CASES)

    # 7: real data table
    s = old[7]
    take_text(s, "TextBox 8")
    texts = take_text(s, "Text real data")
    move(shapes_named(s, "Table real data")[0], top=0.95)
    set_notes(s, "Table 2. The real-data deposits.", *texts)

    # 10 -> 8: comparator methods table
    s = old[10]
    take_text(s, "TextBox 8")
    texts = take_text(s, "Text compared")
    move(shapes_named(s, "Table methods")[0], top=0.95)
    set_notes(s, "Table 3. The comparator methods.", *texts)

    # 11 -> 9: comparator heatmap
    s = old[11]
    texts = take_text(s, "Text methods heatmap")
    drop_pictures(s)
    place(s, FIG / "comparator" / "method_comparison_heatmap.png")
    set_notes(s, *texts, CASES)

    # 12 -> 10: F1 + acceptance table
    s = old[12]
    caption = take_text(s, "Caption acceptance")
    footnote = take_text(s, "Note acceptance")
    texts = take_text(s, "Text f1")
    drop_pictures(s)
    f1 = place(s, FIG / "comparator" / "method_comparison_f1.png", area=(0.35, 0.85, 12.63, 2.6), top_align=True)
    table = shapes_named(s, "Table acceptance")[0]
    move(table, top=f1.top / 914400 + f1.height / 914400 + 0.35, left=(13.333 - table.width / 914400) / 2)
    corner, sample = table.table.cell(0, 0), table.table.cell(0, 1)
    size = next((r.font.size for p in sample.text_frame.paragraphs for r in p.runs if r.font.size), Pt(11))
    run = corner.text_frame.paragraphs[0].add_run()
    run.text = "Acceptance"
    run.font.size = size
    set_notes(s, *texts, "Table: " + " ".join(caption) + ". " + " ".join(footnote))

    # 13 -> 11: rerun table
    s = old[13]
    changes = take_text(s, "Text rerun changes")
    callout = take_text(s, "Callout not yet")
    take_text(s, "TextBox 8")
    caption = take_text(s, "Caption rerun")
    take_text(s, "Text reading")
    move(shapes_named(s, "Table rerun")[0], top=0.95)
    set_notes(s, *changes, *callout, "Table 4. " + " ".join(caption).replace("the top rows of slides 14–16", "slides 14–16"))

    # 14-16 -> 12-14: 32 arms, one dot plot each
    for number, stem in ((14, "ovarian_gse180661"), (15, "colorectal_gse225857"), (16, "prostate_gse271675")):
        s = old[number]
        texts = take_text(s, "Note retained", "Note rejected")
        drop_highlights(s)
        drop_pictures(s)
        place(s, DOT / "gsea_dotplot_factorial" / f"gsea_dotplot_{stem}.png")
        unit = UNIT_NES.format("logcpm_noscale (15 sets) and rank256_ds (19)") if number == 16 else ""
        set_notes(s, DOTPLOT_READING, FACTORIAL_ROWS, unit, *texts)

    # 17 -> 15: cell cycle x CNV design + table 5; old slide 9 goes into its notes
    s = old[17]
    design = take_text(s, "Text cycle arms", "Text representations")
    take_text(s, "TextBox 8")
    caption = take_text(s, "Caption split")
    split = take_text(s, "Text split")
    move(shapes_named(s, "Table split")[0], top=0.95)
    background = [text_of(shape) for shape in old[9].shapes
                  if shape.has_text_frame and shape.name in ("Text infercnv", "Text cell cycle")]
    set_notes(s, *design, "Table 5. " + " ".join(caption), *split,
              "Background (from the 10/02 report's status slide, state as of 2026-10-01):", *background)

    # 18-20 -> 16-18: cycling x CNV dot plots
    for number, stem in ((18, "ovarian_gse180661"), (19, "colorectal_gse225857"), (20, "prostate_gse271675")):
        s = old[number]
        texts = take_text(s, "Notes cycling")
        drop_highlights(s)
        drop_pictures(s)
        place(s, DOT / "gsea_dotplot" / f"gsea_dotplot_{stem}.png")
        texts = [t.replace("Number right of a row: arms with a dot. Strip: cells removed = B, D; genes removed = C, D; "
                           "unmarked label = raw.", "").replace("grey columns were not run", "the RNA+CNV and CNV rows "
                           "were not run").replace("Both rows are framed in orange.", "").strip() for t in texts]
        unit = UNIT_NES.format("CNV arm B (16 sets)") if number == 18 else ""
        set_notes(s, DOTPLOT_READING, CYCLING_ROWS, unit, *texts)

    # 21 -> 19: NES table
    s = old[21]
    caption = take_text(s, "Caption nes")
    texts = take_text(s, "Text nes")
    move(shapes_named(s, "Table nes")[0], top=0.95)
    set_notes(s, "Table: " + " ".join(caption), *texts)

    # 22 -> 20: summary, short on the slide, full text in the notes
    s = old[22]
    full = [text_of(shape) for shape in s.shapes if shape.has_text_frame
            and shape.name in ("Text summary benchmark", "Text summary real data", "Callout open")]
    set_bullets(shapes_named(s, "Text summary benchmark")[0], "Benchmark", [
        "ConfidenceOT: directional F1 0.99 on ranknm256_noscale_ds_cos at acceptance 0.99",
        "Vanilla UOT and IC-POT at their fixed cutoffs reject nothing; Partial OT needs the true m",
        "On raw every method fails (F1 ≤ 0.64)"])
    set_bullets(shapes_named(s, "Text summary real data")[0], "Real data", [
        "E2F and G2M follow the cycling cells; MYC targets V1 stays higher in retained",
        "The split follows the RNA preprocessing; with CNV, ovarian rejection falls from 55–58% to 5–6%",
        "Prostate: EMT and androgen response go with the smaller group"])
    set_bullets(shapes_named(s, "Callout open")[0], "Open", [
        "Acceptance 0.99 was chosen with the truth in view",
        "Prostate CNV arms wait for R inferCNV",
        "Not checked: whether the raw split follows sequencing depth, and whether CNV-only DEGs follow chromosome position"])
    set_notes(s, *full)

    # Drop old slides 8 and 9.
    id_list = prs.slides._sldIdLst
    entries = list(id_list)
    for index in (8, 7):  # zero-based: old slide 9 then 8
        entry = entries[index]
        prs.part.drop_rel(entry.rId)
        id_list.remove(entry)

    for number, s in enumerate(prs.slides, 1):
        frame = s.notes_slide.notes_text_frame
        text = frame.text
        for a, b in FIXES.get(number, []):
            if a not in text:
                raise SystemExit(f"slide {number}: note fix {a!r} not found")
            text = text.replace(a, b)
        frame.text = text

    # Renumber the footer number box.
    for number, s in enumerate(prs.slides, 1):
        # A slide-number field: PowerPoint recomputes it; the cached text is set to match.
        for shape in shapes_named(s, "Rectangle 6"):
            for field in shape._element.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}fld"):
                field.find("{http://schemas.openxmlformats.org/drawingml/2006/main}t").text = str(number)

    prs.save(str(OUT))
    print(OUT, len(prs.slides), "slides")


if __name__ == "__main__":
    main()
