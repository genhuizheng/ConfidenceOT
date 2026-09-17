"""Print the review feedback a .pptx carries, slide by slide.

Three places a reviewer can leave notes, all read here so it does not matter
which one they pick:

* PowerPoint comments. Modern PowerPoint writes one part per comment under
  ``ppt/comments/modernComment_*.xml``; older files put them all in
  ``ppt/comments/commentN.xml`` with the author list in
  ``ppt/commentAuthors.xml``. Both are parsed.
* The speaker-notes pane, which python-pptx exposes directly.
* Text boxes whose content looks like a note rather than slide content, found
  by a marker (default "TODO", "FIX", "??" or a leading "@").

Comments are the best of the three for review: they anchor to a slide and, when
made on a shape, to that shape, and they leave the deck's own content alone.

Usage:
  python cancer_metastasis/tools/read_deck_feedback.py deck.pptx
"""

from __future__ import annotations

import argparse
import io
import re
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree

from pptx import Presentation

# The Windows console defaults to a code page that cannot encode the Chinese
# these notes are written in, and the failure looks like garbled feedback
# rather than a terminal setting.
if hasattr(sys.stdout, "buffer"):
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8",
                                  errors="replace", line_buffering=True)

MAIN = "http://schemas.openxmlformats.org/presentationml/2006/main"
DRAW = "http://schemas.openxmlformats.org/drawingml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
P188 = "http://schemas.microsoft.com/office/powerpoint/2018/8/main"
MARKERS = re.compile(r"(?i)\bTODO\b|\bFIX\b|\?\?|^@")


def slide_order(package: zipfile.ZipFile) -> list[str]:
    """Slide part names in presentation order.

    The zip's own ordering is arbitrary and slideN.xml numbering is creation
    order, not display order, so the order is taken from the presentation's
    slide id list resolved through its relationships.
    """
    presentation = ElementTree.fromstring(package.read("ppt/presentation.xml"))
    rels = ElementTree.fromstring(package.read("ppt/_rels/presentation.xml.rels"))
    targets = {rel.get("Id"): rel.get("Target")
               for rel in rels.findall("{*}Relationship")}
    names = []
    for entry in presentation.iter(f"{{{MAIN}}}sldId"):
        target = targets.get(entry.get(f"{{{REL}}}id"), "")
        if target:
            names.append("ppt/" + target.lstrip("/").replace("../", ""))
    return names


def authors(package: zipfile.ZipFile) -> dict[str, str]:
    names: dict[str, str] = {}
    for part in ("ppt/commentAuthors.xml", "ppt/authors.xml"):
        if part not in package.namelist():
            continue
        root = ElementTree.fromstring(package.read(part))
        for author in root:
            key = author.get("id") or author.get("{%s}id" % P188)
            name = author.get("name") or author.get("initials") or "?"
            if key:
                names[key] = name
    return names


def text_of(element) -> str:
    return "".join(node.text or "" for node in element.iter(f"{{{DRAW}}}t")).strip()


def comments_for(package: zipfile.ZipFile, slide_part: str,
                 people: dict[str, str]) -> list[str]:
    """Every comment attached to one slide, from either storage format."""
    rels_part = slide_part.replace("slides/", "slides/_rels/") + ".rels"
    if rels_part not in package.namelist():
        return []
    rels = ElementTree.fromstring(package.read(rels_part))
    found = []
    for rel in rels.findall("{*}Relationship"):
        target = rel.get("Target", "")
        if "comment" not in target.lower():
            continue
        part = "ppt/" + target.lstrip("/").replace("../", "")
        if part not in package.namelist():
            continue
        root = ElementTree.fromstring(package.read(part))
        # Modern comments carry their text in a:t runs; the legacy format uses
        # a p:text element instead.
        for comment in root.iter():
            tag = comment.tag.split("}")[-1]
            if tag not in ("cm", "cmt"):
                continue
            body = text_of(comment)
            if not body:
                legacy = comment.find(f"{{{MAIN}}}text")
                body = (legacy.text or "").strip() if legacy is not None else ""
            if not body:
                continue
            who = people.get(comment.get("authorId")
                             or comment.get("{%s}authorId" % P188), "?")
            when = (comment.get("dt") or comment.get("created") or "")[:16]
            found.append(f"[{who}{' ' + when if when else ''}] {body}")
    return found


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("deck", type=Path)
    parser.add_argument("--all-text-boxes", action="store_true",
                        help="print every text box, not only marked ones")
    args = parser.parse_args()

    presentation = Presentation(str(args.deck))
    slides = list(presentation.slides)
    with zipfile.ZipFile(args.deck) as package:
        people = authors(package)
        parts = slide_order(package)
        if len(parts) != len(slides):
            parts = [p for p in package.namelist()
                     if re.fullmatch(r"ppt/slides/slide\d+\.xml", p)]
            parts.sort(key=lambda p: int(re.search(r"(\d+)", p).group(1)))
        total = 0
        for index, (slide, part) in enumerate(zip(slides, parts), 1):
            notes = ""
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
            remarks = comments_for(package, part, people)
            marked = []
            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue
                body = shape.text_frame.text.strip()
                if body and (args.all_text_boxes or MARKERS.search(body)):
                    marked.append(body)
            if not (notes or remarks or marked):
                continue
            # The running header and the page number are on every slide, so the
            # first text box is furniture rather than the slide's subject.
            furniture = {"Progress Report", str(index)}
            title = next((s.text_frame.text.strip().splitlines()[0]
                          for s in slide.shapes
                          if s.has_text_frame
                          and s.text_frame.text.strip()
                          and s.text_frame.text.strip() not in furniture), "")
            print(f"\n=== slide {index}  {title[:60]}")
            for remark in remarks:
                print(f"  comment: {remark}")
                total += 1
            if notes:
                print(f"  notes:   {notes}")
                total += 1
            for body in marked:
                print(f"  marked:  {body[:200]}")
                total += 1
        print(f"\n{total} item(s) of feedback"
              + ("" if total else " — nothing found"))


if __name__ == "__main__":
    main()
