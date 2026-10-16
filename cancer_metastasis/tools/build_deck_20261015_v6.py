"""ConfidenceOT_20261015_v6.pptx: the user's edited v5 with slide-size figures on slides 9-18.

The user's v5 (edited by hand, 28 slides) is only read. Every slide keeps its content,
except that each matplotlib figure on slides 9, 10, 12-14 and 16-18 is replaced by the
redraw from slide_figures_20261015_v6.py, drawn at the size it occupies so its text is
read at 9-11 pt. Panels too dense for one slide are split, every row kept:

  9  -> 9, 10    comparator rejected fraction: ConfidenceOT (1/2), comparators (2/2)
  12-14 -> two slides each   GSEA, 32 arms: arms 1-16 (1/2), 17-32 (2/2); all 50 sets on both

A duplicated slide carries its source's title (with the part added) and notes (with one
line naming the part). Slide numbers are renumbered.

    python cancer_metastasis/tools/build_deck_20261015_v6.py
"""

from __future__ import annotations

import importlib.util
import subprocess
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[1]
_spec = importlib.util.spec_from_file_location("deck_v2", HERE / "build_deck_20261015_v2.py")
v2 = importlib.util.module_from_spec(_spec)
sys.modules["deck_v2"] = v2
_spec.loader.exec_module(v2)

SRC = Path(r"D:\Xia_lab\Progress_report\progress_report7\ConfidenceOT_20261015_v5.pptx")
OUT = Path(r"D:\Xia_lab\Progress_report\progress_report7\ConfidenceOT_20261015_v6.pptx")
TMP = OUT.with_name("_tmp_v6.pptx")
FIG = REPO / "figures" / "slides_20261015_v6"
AREA = (0.35, 0.85, 12.63, 6.25)  # x, y, w, h below the title, inches
SPLIT = {"slide9.xml": ("comparator", "Comparator benchmark: rejected fraction"),
         "slide12.xml": ("ovarian", "Ovarian: all hallmark sets, 32 arms"),
         "slide13.xml": ("colorectal", "Colorectal: all hallmark sets, 32 arms"),
         "slide14.xml": ("prostate", "Prostate: all hallmark sets, 32 arms")}
CYCLING = {"slide16.xml": "ovarian", "slide17.xml": "colorectal", "slide18.xml": "prostate"}


def title_of(slide) -> str:
    return next(s for s in slide.shapes if s.name == "TextBox 3").text_frame.text


def figure_pictures(slide):
    return [s for s in slide.shapes if s.shape_type == 13 and s.name != "Picture 8"]


def place(slide, path: Path, area=AREA, name="Figure") -> None:
    for old in figure_pictures(slide):
        old._element.getparent().remove(old._element)
    picture = slide.shapes.add_picture(str(path), Inches(area[0]), Inches(area[1]))
    x, y, w, h = area
    scale = min(Inches(w) / picture.width, Inches(h) / picture.height)
    picture.width, picture.height = int(picture.width * scale), int(picture.height * scale)
    picture.left = int(Inches(x) + (Inches(w) - picture.width) / 2)
    picture.top = int(Inches(y))
    picture.name = name


def copy_notes(source, target, extra: str) -> None:
    text = source.notes_slide.notes_text_frame.text if source.has_notes_slide else ""
    target.notes_slide.notes_text_frame.text = (extra + "\n" + text).strip()


def add_note_line(slide, extra: str) -> None:
    frame = slide.notes_slide.notes_text_frame
    frame.text = (extra + "\n" + frame.text).strip()


def main() -> None:
    script = v2.SKILL / "scripts" / "add_slide.py"
    current = SRC
    for part in SPLIT:
        subprocess.run([sys.executable, str(script), str(current), part, "--after", part, "-o", str(TMP)],
                       check=True, capture_output=True, text=True, encoding="utf-8")
        current = TMP
    prs = Presentation(str(TMP))
    slides = list(prs.slides)
    assert len(slides) == 32, len(slides)
    # python-pptx renames slide parts by position on load, so slides are found by title. Each split
    # slide's duplicate sits right after it, with the same title.
    def find(title: str) -> list:
        return [s for s in slides if title_of(s) == title]
    by_part, duplicates = {}, {}
    for part, (_, base) in SPLIT.items():
        pair = find(base)
        assert len(pair) == 2 and slides.index(pair[1]) == slides.index(pair[0]) + 1, (base, len(pair))
        by_part[part], duplicates[part] = pair
    for part, base in (("slide10.xml", "Comparator benchmark: directional F1"),
                       ("slide16.xml", "Ovarian: all hallmark sets, cell cycle × CNV"),
                       ("slide17.xml", "Colorectal: all hallmark sets, cell cycle × CNV"),
                       ("slide18.xml", "Prostate: all hallmark sets, cell cycle × CNV")):
        found = find(base)
        assert len(found) == 1, (base, len(found))
        by_part[part] = found[0]

    source_note = ("Figure redrawn at slide size (cancer_metastasis/tools/slide_figures_20261015_v6.py) from the "
                   "original figure's value table; values, colour scales and encodings unchanged.")
    # Comparator heatmap, split by method family.
    first, second = by_part["slide9.xml"], duplicates["slide9.xml"]
    _, base = SPLIT["slide9.xml"]
    v2.set_title(first, base + " (1/2)")
    v2.set_title(second, base + " (2/2)")
    place(first, FIG / "comparator_heatmap_confidenceot.png")
    place(second, FIG / "comparator_heatmap_comparators.png")
    copy_notes(first, second, "Part 2 of 2: the comparators (UOT, Partial OT, IC-POT, balanced OT); same colour "
                              "scale and columns as part 1. " + source_note)
    add_note_line(first, "Part 1 of 2: ConfidenceOT M4-E and M4-R at acceptance 0.90, 0.95, 0.99. " + source_note)

    # Directional F1: above the table, which stays where the user put it.
    f1_slide = by_part["slide10.xml"]
    assert title_of(f1_slide).startswith("Comparator benchmark: directional F1")
    table = next(s for s in f1_slide.shapes if s.name == "Table acceptance")
    space = table.top / 914400 - 0.08 - AREA[1]
    place(f1_slide, FIG / "comparator_f1.png", area=(AREA[0], AREA[1], AREA[2], space))
    add_note_line(f1_slide, source_note)

    # GSEA, 32 arms: two slides each.
    for part in ("slide12.xml", "slide13.xml", "slide14.xml"):
        dataset, base = SPLIT[part]
        first, second = by_part[part], duplicates[part]
        v2.set_title(first, base + " (1/2)")
        v2.set_title(second, base + " (2/2)")
        place(first, FIG / f"gsea_factorial_{dataset}_part1.png")
        place(second, FIG / f"gsea_factorial_{dataset}_part2.png")
        copy_notes(first, second, "Part 2 of 2: arms 17-32 (rank256 and ranknm256 families); all 50 hallmark sets. "
                                  + source_note)
        add_note_line(first, "Part 1 of 2: arms 1-16 (raw and logcpm families); all 50 hallmark sets. " + source_note)

    # GSEA, cycling x CNV: one slide each.
    for part, dataset in CYCLING.items():
        place(by_part[part], FIG / f"gsea_cycling_{dataset}.png")
        add_note_line(by_part[part], source_note)

    for number, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if shape.name == "Rectangle 6":
                for field in shape._element.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}fld"):
                    field.find("{http://schemas.openxmlformats.org/drawingml/2006/main}t").text = str(number)
    prs.save(str(OUT))
    TMP.unlink()
    print(OUT, len(prs.slides), "slides")


if __name__ == "__main__":
    main()
