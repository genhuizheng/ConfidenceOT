"""ConfidenceOT_20261015_v4.pptx: v3 with the tumour-level section rebuilt around the ranking task.

v3 is only read. The KP-Tracer part becomes:

20 KP-Tracer: Independent Lineage Evidence                      (kept)
21 How CRISPR Lineage Tracing Works                             (kept)
22 Tumor-Level Benchmark: Inputs, Outputs, and What We Evaluate  (new design; replaces the v3 ranking schematic)
23 Retained Fractions and Lineage Ground Truth (1/2)            (heatmap with a per-lesion rank annotation)
24 Retained Fractions and Lineage Ground Truth (2/2)
25 Does ConfidenceOT Rank the Related Primary First?            (metric definitions + mouse-level dot plot)
26 RNA Preprocessing and Family Discrimination                  (kept)
27 Raw vs Preprocessed RNA: Matched Heatmaps (1/2)              (with the rank annotation)
28 Raw vs Preprocessed RNA: Matched Heatmaps (2/2)
29 CNV and RNA+CNV: Not Yet Evaluated                           (kept)
30 What the Tumor-Level Benchmark Supports                      (revised)
31 Summary and next steps                                       (kept)

The v3 lesion-level d_M bar chart and mouse-level d_M plot are dropped as slides; the
prespecified d_M result stays as a supporting line.

Rank annotation per lesion row: the rank of the best-ranked same-family candidate among the
lesion's completed candidates, by retained fraction: "1/3" means it is first of three.
Hit = ranked first with no other-family candidate tied with it; "tie" when an other-family
candidate has exactly the same r; "—" when the lesion lacks a same-family or an
other-family candidate (no ranking question).
"""
import importlib.util
import json
import subprocess
import sys
from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_PATTERN
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
_spec = importlib.util.spec_from_file_location("deck_v3", HERE / "build_deck_20261015_v3.py")
v3 = importlib.util.module_from_spec(_spec)
sys.modules["deck_v3"] = v3
_spec.loader.exec_module(v3)
v2 = v3.v2
text, box, arrow, segment, dots, notes = v2.text, v2.box, v2.arrow, v2.segment, v2.dots, v2.notes
ORANGE, FILL, INK, MUTED, GREY, PALE = v2.ORANGE, v2.FILL, v2.INK, v2.MUTED, v2.GREY, v2.PALE
RETAINED, REJECTED, FONT, WHITE = v2.RETAINED, v2.REJECTED, v2.FONT, v3.WHITE
DARK_GREY = RGBColor(0x7F, 0x7F, 0x7F)
LIGHT_LINE = RGBColor(0xE6, 0xE6, 0xE6)
colour, short = v3.colour, v3.short

SRC = Path(r"D:\Xia_lab\Progress_report\progress_report7\ConfidenceOT_20261015_v3.pptx")
OUT = Path(r"D:\Xia_lab\Progress_report\progress_report7\ConfidenceOT_20261015_v4.pptx")
TMP = OUT.with_name("_tmp_v4.pptx")
EXPLORE, MAIN, RAW = v3.EXPLORE, v3.MAIN, v3.RAW


# --- rank of the best same-family candidate ------------------------------------------------

def best_rank(block: dict, lesion: str):
    """(rank, n, ties) for a lesion, or None when it has no ranking question."""
    values = [(v, fam) for (les, _), (v, fam) in block["cells"].items() if les == lesion and v is not None]
    positives = [v for v, fam in values if fam]
    negatives = [v for v, fam in values if not fam]
    if not positives or not negatives:
        return None
    best = max(positives)
    above = sum(1 for v in negatives if v > best)
    tied = sum(1 for v in negatives if v == best)
    return above + 1, len(values), tied


def annotation(block: dict, lesion: str):
    result = best_rank(block, lesion)
    if result is None:
        return [("—", {"color": GREY})]
    rank, n, tied = result
    label = f"{rank}/{n}" if not tied else f"{rank}–{rank + tied}/{n}"
    if rank == 1 and not tied:
        return [(label, {"bold": True}), ("  hit", {"bold": True, "color": ORANGE})]
    if tied:
        return [(label, {}), ("  tie", {"color": MUTED})]
    return [(label, {}), ("  miss", {"color": MUTED})]


def draw_block(slide, mouse: str, data: dict, x: float, y: float, cell: float, label_w: float, annot_w: float,
               font: int = 11):
    block = data[mouse]
    grid_w = cell * len(block["primaries"])
    text(slide, x, y, label_w + grid_w + annot_w, 0.3,
         [[(f"mouse {mouse}", {"bold": True}), (f"  {block['genotype']}", {"color": MUTED})]], size=font + 1)
    top = y + 0.3
    for c, primary in enumerate(block["primaries"]):
        text(slide, x + label_w + c * cell, top, cell, 0.26, short(primary), size=font - 1, color=MUTED,
             align=PP_ALIGN.CENTER)
    text(slide, x + label_w + grid_w + 0.06, top, annot_w, 0.26, "GT rank", size=font - 2, color=MUTED)
    grid_top = top + 0.28
    for r, lesion in enumerate(block["lesions"]):
        yy = grid_top + r * cell
        informative = block["informative"][lesion]
        text(slide, x, yy, label_w - 0.05, cell, short(lesion), size=font - 1, color=INK if informative else GREY,
             italic=not informative, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
        for c, primary in enumerate(block["primaries"]):
            value, family = block["cells"][(lesion, primary)]
            square = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x + label_w + c * cell + 0.02),
                                            Inches(yy + 0.02), Inches(cell - 0.04), Inches(cell - 0.04))
            square.shadow.inherit = False
            if value is None:
                square.fill.patterned()
                square.fill.pattern = MSO_PATTERN.LIGHT_UPWARD_DIAGONAL
                square.fill.fore_color.rgb = GREY
                square.fill.back_color.rgb = WHITE
            else:
                square.fill.solid()
                square.fill.fore_color.rgb = colour(value)
            if family:
                square.line.color.rgb = INK
                square.line.width = Pt(2.5)
            else:
                square.line.fill.background()
        text(slide, x + label_w + grid_w + 0.06, yy, annot_w, cell, [annotation(block, lesion)], size=font - 1,
             anchor=MSO_ANCHOR.MIDDLE)
    return label_w + grid_w + annot_w, 0.58 + cell * len(block["lesions"])


def draw_rows(slide, data, rows_of_mice, x0, y0, cell, label_w, annot_w, gap=0.35, row_gap=0.25, font=11):
    y = y0
    for mice in rows_of_mice:
        x, height = x0, 0.0
        for mouse in mice:
            w, h = draw_block(slide, mouse, data, x, y, cell, label_w, annot_w, font)
            x += w + gap
            height = max(height, h)
        y += height + row_gap
    return y


def legend(slide, x, y):
    v3.draw_legend(slide, x, y)
    yy = y + 2.25
    text(slide, x, yy, 2.9, 1.5, [[("GT rank", {"bold": True})],
                                  [("1/3 hit", {"bold": True}), ("  the best same-family primary is ranked 1st of 3",
                                                                  {"color": MUTED})],
                                  [("2/4 miss", {}), ("  ranked 2nd of 4", {"color": MUTED})],
                                  [("—", {}), ("  no ranking question", {"color": MUTED})]], size=11)


def tally(data: dict) -> tuple[int, int]:
    hits = total = 0
    for block in data.values():
        for lesion in block["lesions"]:
            result = best_rank(block, lesion)
            if result is None:
                continue
            total += 1
            hits += result[0] == 1 and result[2] == 0
    return hits, total


# --- slides ----------------------------------------------------------------------------------

def slide_task(slide) -> None:
    heads = [("Given", 0.45), ("Model input", 3.6), ("Model output", 6.75), ("Independent ground truth", 9.9)]
    for head, x in heads:
        text(slide, x, 0.95, 3.0, 0.35, head, size=15, bold=True, color=ORANGE)
    for _, x in heads[1:]:
        arrow(slide, x - 0.3, 2.35, x - 0.08, 2.35)
    rows = (("P1", 0.70, True), ("P2", 0.40, False), ("P3", 0.30, False))

    # Given: one metastasis, three candidate primaries.
    x = 0.45
    box(slide, x + 1.75, 1.55, 1.1, 1.3, None, line=ORANGE, name="M")
    text(slide, x + 1.8, 1.58, 0.5, 0.3, "M", size=13, bold=True)
    dots(slide, x + 1.88, 1.95, 4, 4, [GREY] * 16, size=0.12, gap=0.07)
    for k, (name, _, _) in enumerate(rows):
        box(slide, x, 1.45 + k * 0.6, 0.95, 0.45, name, line=ORANGE, size=13, bold=True)
    text(slide, x, 3.35, 3.0, 1.0, ["one metastatic lesion M", "candidate primaries P1–P3 from the same mouse"],
         size=12)

    # Input.
    x = 3.6
    box(slide, x, 1.45, 2.85, 1.75, ["one ConfidenceOT run per pair", "source: primary malignant cells",
                                     "target: metastatic malignant cells"], fill=FILL, line=ORANGE, size=13, bold=True)
    text(slide, x, 3.35, 3.0, 1.0, [[("RNA expression only", {"bold": True})], "P1 → M,  P2 → M,  P3 → M"], size=12)

    # Output: one retained fraction per pair.
    x = 6.75
    for k, (name, r, _) in enumerate(rows):
        y = 1.5 + k * 0.6
        text(slide, x, y, 0.45, 0.4, name, size=13, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x + 0.5), Inches(y + 0.06), Inches(2.0 * r),
                                     Inches(0.3))
        bar.fill.solid()
        bar.fill.fore_color.rgb = colour(r)
        bar.line.fill.background()
        bar.shadow.inherit = False
        text(slide, x + 0.56 + 2.0 * r, y + 0.03, 0.6, 0.35, f"{r:.2f}", size=12, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
    text(slide, x, 3.35, 3.0, 1.0, [[("r(P, M)", {"bold": True}), (": the share of P's cells retained as having a "
                                                                    "counterpart in M", {})]], size=12)

    # Ground truth.
    x = 9.9
    for k, (name, _, same) in enumerate(rows):
        y = 1.5 + k * 0.6
        text(slide, x, y, 0.45, 0.4, name, size=13, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        box(slide, x + 0.55, y + 0.03, 1.6, 0.36, "same family" if same else "other family",
            fill=FILL if same else None, line=ORANGE, dashed=True, size=12)
    text(slide, x, 3.35, 3.0, 1.0, ["from the CRISPR lineage data (authors' MetFamily)",
                                    [("never seen by ConfidenceOT", {"bold": True, "color": REJECTED})]], size=12)

    # What is evaluated.
    box(slide, 0.45, 4.5, 12.4, 1.75, None, fill=PALE, line=None, name="Evaluate")
    text(slide, 0.65, 4.58, 3.0, 0.35, "What we evaluate", size=15, bold=True)
    text(slide, 0.65, 5.0, 6.0, 1.2, [[("1  ", {"bold": True, "color": ORANGE}),
                                       ("Do same-family primaries have higher retained fractions?", {"bold": True})],
                                      [("    prespecified: d", {"color": MUTED}), ("M", {"color": MUTED, "baseline": -25000}),
                                       (" = mean r(same) − mean r(other)", {"color": MUTED})]], size=13)
    text(slide, 6.85, 5.0, 5.9, 1.2, [[("2  ", {"bold": True, "color": ORANGE}),
                                       ("Do they rank higher among the candidates?", {"bold": True})],
                                      [("    here P1 is ranked 1st of 3: a hit", {"color": MUTED})],
                                      [("    exploratory: top-1 hit, reciprocal rank, pairwise AUROC", {"color": MUTED})]],
         size=13)
    text(slide, 0.45, 6.55, 12.4, 0.35, "Schematic: the cells and values are illustrative, not data.", size=11,
         color=MUTED, italic=True)
    notes(slide, "\n".join([
        "The decision problem first. Given one metastatic lesion M and the primary tumours of the same mouse, which "
        "primary does M belong with?",
        "Input: for each candidate primary P, one ConfidenceOT run with P's malignant cells as source and M's "
        "malignant cells as target, from RNA expression only (ranknm256_noscale_ds_cos, M4-E, acceptance 0.99). "
        "Malignant = not is_normal, the authors' call; pairs need at least 20 malignant cells per side.",
        "Output: one number per pair, r(P, M), the share of P's analysed cells the gate retains as having a "
        "counterpart in M.",
        "Ground truth: the authors' MetFamily labels from the lineage data mark each candidate as same-family or "
        "other-family. Several candidates may be same-family. The labels are attached only after the runs.",
        "Two questions: (1) do same-family candidates have higher r? This is the prespecified readout d_M (27 of 31 "
        "lesions > 0; one-sided Wilcoxon p = 4.99e-7; 8 of 8 mice positive). (2) Do they rank higher among the "
        "candidates? Exploratory ranking metrics, slide 25.",
        "Not implied: that any sampled primary cell seeded M, or the route of spread.",
    ]))


def slide_heatmap(slide, data: dict, rows_of_mice, part: str, hits: tuple) -> None:
    draw_rows(slide, data, rows_of_mice, 0.45, 0.95, 0.34, 0.5, 0.95)
    legend(slide, 10.2, 1.0)
    text(slide, 10.2, 4.75, 2.8, 1.6, [[(f"{hits[0]} of {hits[1]}", {"bold": True, "size": 22, "color": ORANGE})],
                                       [("lesions where a same-family primary ranks first (all mice, ties not "
                                         "counted)", {"color": MUTED})]], size=12)
    text(slide, 0.45, 6.75, 12.4, 0.3, f"ranknm256_noscale_ds_cos, M4-E; {part}. Source: readout/pairs_{MAIN}.csv; "
         "primaries in tumour order within each mouse; one colour scale for every mouse.", size=10, color=MUTED,
         italic=True)
    notes(slide, "\n".join([
        "Real data. Each block is one mouse; rows are its metastatic lesions, columns its candidate primaries in tumour "
        "order (not ordered by the labels). Colour: retained fraction r on one 0–1 scale for every mouse. Black "
        "outline: same family as the lesion (lineage ground truth). Hatched: not run (fewer than 20 malignant cells "
        "on a side), which is not the same as r = 0.",
        "GT rank: the rank of the best-ranked same-family primary among the lesion's completed candidates, by r. "
        "1/3 hit = ranked first of three; 2/4 miss = second of four; '—' = the lesion lacks a same-family or an "
        "other-family candidate, so there is no ranking question (grey italic lesion).",
        f"Across both heatmap slides: a same-family primary ranks first in {hits[0]} of {hits[1]} lesions with a "
        "ranking question (strict: an exact tie with an other-family candidate is not counted as a hit).",
        "Exploratory reading; the prespecified readout is d_M.",
    ]))


def dot_plot(slide, mouse_table: pd.DataFrame, x0, y0, width, height) -> None:
    metrics = (("top1", "Top-1 hit"), ("rr", "Reciprocal rank"), ("auroc", "Pairwise AUROC"))
    series = ((MAIN, "ranknm256_noscale_ds_cos", ORANGE, ORANGE, False), (RAW, "raw", DARK_GREY, DARK_GREY, False),
              (None, "random ranking", WHITE, DARK_GREY, True))
    left, bottom, top = x0 + 0.55, y0 + height - 0.55, y0 + 0.45
    right = x0 + width
    span = bottom - top

    def ypos(v):
        return bottom - span * v
    for tick in (0, 0.25, 0.5, 0.75, 1.0):
        segment(slide, left, ypos(tick), right, ypos(tick), color=LIGHT_LINE if tick else INK, width=0.75 if tick else 1.25)
        text(slide, x0, ypos(tick) - 0.13, 0.48, 0.26, f"{tick:g}", size=11, color=MUTED, align=PP_ALIGN.RIGHT)
    column = (right - left) / len(metrics)
    mice = sorted(mouse_table["patient_id"].astype(str).unique(), key=int)
    for k, (metric, name) in enumerate(metrics):
        centre = left + column * (k + 0.5)
        text(slide, centre - column / 2, bottom + 0.08, column, 0.3, name, size=13, align=PP_ALIGN.CENTER)
        for s, (label, _, fill, line, hollow) in enumerate(series):
            base = centre + (s - 1) * column * 0.27
            table = mouse_table[mouse_table["label"].eq(label if label else MAIN)]
            table = table.set_index(table["patient_id"].astype(str))
            column_name = metric if label else f"{metric}_random"
            values = []
            for i, mouse in enumerate(mice):
                value = float(table.loc[mouse, column_name])
                values.append(value)
                jitter = (i - (len(mice) - 1) / 2) * 0.035
                dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(base + jitter - 0.065), Inches(ypos(value) - 0.065),
                                             Inches(0.13), Inches(0.13))
                dot.shadow.inherit = False
                dot.fill.solid()
                dot.fill.fore_color.rgb = fill
                dot.line.color.rgb = line
                dot.line.width = Pt(1.0)
            mean = sum(values) / len(values)
            segment(slide, base - 0.22, ypos(mean), base + 0.22, ypos(mean), color=INK, width=2.5)
    # Legend.
    lx = left
    for label, name, fill, line, hollow in series:
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(lx), Inches(y0 + 0.08), Inches(0.13), Inches(0.13))
        dot.shadow.inherit = False
        dot.fill.solid()
        dot.fill.fore_color.rgb = fill
        dot.line.color.rgb = line
        text(slide, lx + 0.18, y0, 2.6, 0.3, name, size=12, color=MUTED)
        lx += 2.45 if label == MAIN else 1.0 if label == RAW else 0
    segment(slide, lx + 1.6, y0 + 0.145, lx + 1.95, y0 + 0.145, color=INK, width=2.5)
    text(slide, lx + 2.0, y0, 1.6, 0.3, "mean of mice", size=12, color=MUTED)


def slide_ranking(slide, ranking: dict, mouse_table: pd.DataFrame, hits: tuple) -> None:
    box(slide, 0.45, 0.95, 4.6, 5.45, None, fill=PALE, line=None, name="Definitions")
    text(slide, 0.62, 1.05, 4.3, 5.3, [
        [("Per lesion: rank its candidate primaries by r", {"bold": True})],
        "",
        [("Top-1 hit", {"bold": True, "color": ORANGE}), (": a same-family primary is ranked first (1 or 0)", {})],
        [("Reciprocal rank", {"bold": True, "color": ORANGE}),
         (": 1 / rank of the best same-family primary (1, ½, ⅓ …)", {})],
        [("Pairwise AUROC", {"bold": True, "color": ORANGE}),
         (": the probability that a same-family primary has a higher r than an other-family one", {})],
        "",
        [("Random ranking", {"bold": True}), (": the value expected from a random order; it depends on the numbers of "
                                              "candidates and same-family candidates (1 of 3 → top-1 = ⅓)", {})],
        "",
        [("Each dot is one mouse: the mean over its lesions. Exploratory, computed after the prespecified d", {"color": MUTED}),
         ("M", {"color": MUTED, "baseline": -25000}), (" result; ties split at random", {"color": MUTED})],
    ], size=13)
    dot_plot(slide, mouse_table, 5.35, 0.95, 7.5, 5.15)
    main, raw = ranking["labels"][MAIN], ranking["labels"][RAW]
    text(slide, 5.35, 6.12, 7.5, 0.4, [[(f"above random in {main['top1']['mice_above_random']}/8 mice", {"bold": True, "color": ORANGE}),
                                        ("  (ranknm256_noscale_ds_cos, all three)", {"color": MUTED}),
                                        (f"   {raw['top1']['mice_above_random']}/8", {"bold": True, "color": DARK_GREY}),
                                        ("  (raw)", {"color": MUTED})]], size=13, align=PP_ALIGN.CENTER)
    text(slide, 0.45, 6.55, 12.4, 0.3, [[("Supporting, prespecified: d", {}), ("M", {"baseline": -25000}),
                                         (" > 0 in 27 of 31 lesions, one-sided Wilcoxon p = 5.0 × 10", {}),
                                         ("−7", {"baseline": 30000}), ("; median d", {}), ("M", {"baseline": -25000}),
                                         (" > 0 in 8 of 8 mice.", {})]], size=12, color=MUTED)
    text(slide, 0.45, 6.8, 12.4, 0.25, "Source: tumour_level_exploratory/mouse_ranking.tsv, ranking_summary.json; "
         "31 lesions in 8 mice.", size=10, color=MUTED, italic=True)
    notes(slide, "\n".join([
        "Question: given a metastasis, does ConfidenceOT rank the lineage-related primary tumours higher?",
        "For each of the 31 lesions with both same-family and other-family candidates, the candidates are ranked by "
        "their retained fraction r. Top-1 hit = a same-family primary is first. Reciprocal rank = 1 / rank of the "
        "best-ranked same-family primary. Pairwise AUROC = the share of same-family vs other-family candidate pairs in "
        "which the same-family one has the higher r. Ties are split at random (exact expectation).",
        "Random ranking: the exact expectation for each lesion's numbers of candidates n and same-family candidates p "
        "(top-1 = p/n; AUROC = 0.5), averaged the same way. It differs between mice because their lesions have "
        "different numbers of candidates.",
        "Each dot is one mouse (the mean over its lesions); the bar is the mean over the 8 mice. Mice are the "
        "independent unit; lesions within a mouse share their primaries.",
        f"ranknm256_noscale_ds_cos: top-1 {main['top1']['mean_of_mouse_means']:.2f} (random "
        f"{main['top1']['random_mean_of_mouse_means']:.2f}; 95% bootstrap over mice {main['top1']['bootstrap_95'][0]:.2f}–"
        f"{main['top1']['bootstrap_95'][1]:.2f}), reciprocal rank {main['rr']['mean_of_mouse_means']:.2f} (random "
        f"{main['rr']['random_mean_of_mouse_means']:.2f}), AUROC {main['auroc']['mean_of_mouse_means']:.2f} (0.5); above "
        f"random in {main['top1']['mice_above_random']} of 8 mice on each.",
        f"raw: top-1 {raw['top1']['mean_of_mouse_means']:.2f}, reciprocal rank {raw['rr']['mean_of_mouse_means']:.2f}, "
        f"AUROC {raw['auroc']['mean_of_mouse_means']:.2f}; above random in {raw['top1']['mice_above_random']} of 8 mice.",
        f"Strict count on the heatmaps: a same-family primary ranks first in {hits[0]} of {hits[1]} lesions.",
        "Exploratory: these metrics were defined after the prespecified d_M result was known, and AUROC uses the same "
        "scores as d_M, so they are not independent confirmation.",
    ]))


def slide_matched(slide, data_main, data_raw, mice: list, side_by_side: bool) -> None:
    if side_by_side:
        text(slide, 0.45, 0.9, 6.0, 0.3, "ranknm256_noscale_ds_cos", size=14, bold=True, color=ORANGE)
        draw_rows(slide, data_main, [mice], 0.45, 1.25, 0.26, 0.42, 0.85, gap=0.22, font=10)
        text(slide, 6.85, 0.9, 6.0, 0.3, "raw", size=14, bold=True, color=DARK_GREY)
        draw_rows(slide, data_raw, [mice], 6.85, 1.25, 0.26, 0.42, 0.85, gap=0.22, font=10)
    else:
        text(slide, 0.45, 0.9, 6.0, 0.3, "ranknm256_noscale_ds_cos", size=14, bold=True, color=ORANGE)
        bottom = draw_rows(slide, data_main, [mice], 0.45, 1.2, 0.28, 0.45, 0.9, gap=0.35, font=10)
        text(slide, 0.45, bottom - 0.05, 6.0, 0.3, "raw", size=14, bold=True, color=DARK_GREY)
        draw_rows(slide, data_raw, [mice], 0.45, bottom + 0.25, 0.28, 0.45, 0.9, gap=0.35, font=10)
    text(slide, 0.45, 6.75, 12.4, 0.3, "Same lesions, candidate order, colour scale, outlines and GT-rank rule as slides "
         "23–24. Source: readout/pairs_*.csv.", size=10, color=MUTED, italic=True)
    notes(slide, "\n".join([
        "The same heatmap and GT-rank annotation for the two representations, on the 8 mice whose lesions have a "
        "ranking question, with identical ordering, colour scale and markers.",
        "raw retains nearly every source cell in nearly every pair (median r 0.99), so its rows are uniformly dark and "
        "the ranking is decided by small differences; ranknm256_noscale_ds_cos retains a median 0.60.",
        "The two representations differ in several preprocessing settings at once (slide 26).",
    ]))


def slide_summary(slide, ranking: dict, comparison: dict, hits: tuple) -> None:
    main = ranking["labels"][MAIN]
    columns = (("Prespecified, supported", ORANGE, [
        [("Same-family primaries retain more source cells: d", {}), ("M", {"baseline": -25000}),
         (" > 0 in 27 of 31 lesions, one-sided Wilcoxon p = 5.0 × 10", {}), ("−7", {"baseline": 30000}),
         ("; median d", {}), ("M", {"baseline": -25000}), (" > 0 in 8 of 8 mice", {})],
        "Lineage labels never entered ConfidenceOT",
    ]), ("Exploratory", MUTED, [
        f"Given a metastasis, a same-family primary ranks first in {hits[0]} of {hits[1]} lesions",
        f"Top-1 {main['top1']['mean_of_mouse_means']:.2f} vs random {main['top1']['random_mean_of_mouse_means']:.2f}; "
        f"above random in {main['top1']['mice_above_random']} of 8 mice",
        f"raw is weaker (5 of 8 mice) but retains a median {comparison['source_retained_fraction_median'][RAW]:.2f} of "
        "cells and differs in several steps",
    ]), ("Not shown here", REJECTED, [
        "Which primary cells seeded a metastasis, or the route",
        "Cell-level lineage concordance: E1 inconclusive; Level 3 coupling analysis not yet run",
        "One-step RNA preprocessing, CNV and RNA+CNV on KP-Tracer: not run",
    ]))
    for k, (head, colour_value, items) in enumerate(columns):
        x = 0.45 + k * 4.2
        box(slide, x, 1.0, 3.95, 3.6, None, fill=PALE if k else FILL, line=None)
        text(slide, x + 0.2, 1.15, 3.6, 0.4, head, size=17, bold=True, color=colour_value)
        text(slide, x + 0.2, 1.7, 3.55, 2.8, [item for pair in zip(items, [""] * len(items)) for item in pair][:-1],
             size=13)
    text(slide, 0.45, 6.55, 12.4, 0.35, "Tumour level: ranknm256_noscale_ds_cos unless stated; mouse = independent unit.",
         size=11, color=MUTED, italic=True)
    notes(slide, "\n".join([
        "What the KP-Tracer tumour-level benchmark supports, as of 2026-10-15.",
        "Prespecified and supported: d_M > 0 across lesions (27 of 31; one-sided Wilcoxon p = 4.99e-7); every mouse "
        "positive (the mouse-level sign test, p = 0.0039, is exploratory).",
        f"Exploratory: the ranking view of the same scores; a same-family primary ranks first in {hits[0]} of {hits[1]} "
        "lesions; top-1, reciprocal rank and AUROC above random in 8 of 8 mice; raw weaker.",
        "Not shown: seeding cells or routes; cell-level lineage concordance (E1 inconclusive; Level 3 not yet run); "
        "other preprocessing, CNV and RNA+CNV on KP-Tracer.",
    ]))


def main():
    ranking = json.loads((EXPLORE / "ranking_summary.json").read_text())
    comparison = ranking["representation_comparison"]
    mouse_table = pd.read_csv(EXPLORE / "mouse_ranking.tsv", sep="\t")
    data_main, data_raw = v3.heatmap_data(MAIN), v3.heatmap_data(RAW)
    hits = tally(data_main)

    script = v2.SKILL / "scripts" / "add_slide.py"
    subprocess.run([sys.executable, str(script), str(SRC), "slide13.xml", "-o", str(TMP)],
                   check=True, capture_output=True, text=True, encoding="utf-8")
    subprocess.run([sys.executable, str(script), str(TMP), "slide13.xml"],
                   check=True, capture_output=True, text=True, encoding="utf-8")
    prs = Presentation(str(TMP))
    slides = list(prs.slides)
    old, new = slides[:31], slides[31:]
    titles = [next(s for s in sl.shapes if s.name == "TextBox 3").text_frame.text for sl in old]
    assert titles[22].startswith("Retained Fractions") and titles[23].startswith("Tumor-Level Concordance")
    assert titles[25].startswith("Candidate-Primary Ranking") and titles[27].startswith("Raw vs Preprocessed")

    plan = [
        (old[21], "Tumor-Level Benchmark: Inputs, Outputs, and What We Evaluate", lambda s: slide_task(s)),
        (old[22], "Retained Fractions and Lineage Ground Truth (1/2)",
         lambda s: slide_heatmap(s, data_main, [["3451", "3454", "3457"], ["3460", "3465", "3466"]], "part 1 of 2", hits)),
        (new[0], "Retained Fractions and Lineage Ground Truth (2/2)",
         lambda s: slide_heatmap(s, data_main, [["3467", "3508", "3515"], ["3513", "3724"]], "part 2 of 2", hits)),
        (old[25], "Does ConfidenceOT Rank the Related Primary First?",
         lambda s: slide_ranking(s, ranking, mouse_table, hits)),
        (old[27], "Raw vs Preprocessed RNA: Matched Heatmaps (1/2)",
         lambda s: slide_matched(s, data_main, data_raw, ["3451", "3454", "3457", "3466", "3513"], False)),
        (new[1], "Raw vs Preprocessed RNA: Matched Heatmaps (2/2)",
         lambda s: slide_matched(s, data_main, data_raw, ["3460", "3508", "3515"], True)),
        (old[29], "What the Tumor-Level Benchmark Supports", lambda s: slide_summary(s, ranking, comparison, hits)),
    ]
    for slide, title, builder in plan:
        v2.clear_content(slide)
        v2.set_title(slide, title)
        if len(title) > 55:  # keep long titles on one line: widen leftwards (the logo ends at 3.4 in) and shrink
            title_box = next(s for s in slide.shapes if s.name == "TextBox 3")
            right_edge = title_box.left + title_box.width
            title_box.left = Inches(3.65)
            title_box.width = right_edge - title_box.left
            for run in title_box.text_frame.paragraphs[0].runs:
                run.font.size = Pt(26)
        builder(slide)

    id_list = prs.slides._sldIdLst
    entries = list(id_list)
    by_slide = dict(zip([id(s) for s in slides], entries))
    drop = [old[23], old[24]]  # the v3 standalone d_M slides
    order = old[:23] + [new[0], old[25], old[26], old[27], new[1], old[28], old[29], old[30]]
    for entry in entries:
        id_list.remove(entry)
    for slide in order:
        id_list.append(by_slide[id(slide)])
    for slide in drop:
        prs.part.drop_rel(by_slide[id(slide)].rId)
    for number, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if shape.name == "Rectangle 6":
                for field in shape._element.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}fld"):
                    field.find("{http://schemas.openxmlformats.org/drawingml/2006/main}t").text = str(number)
    prs.save(str(OUT))
    TMP.unlink()
    print(OUT, len(prs.slides), "slides; strict hits", hits)


if __name__ == "__main__":
    main()
