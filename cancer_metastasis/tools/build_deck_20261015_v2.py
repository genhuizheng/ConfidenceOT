"""ConfidenceOT_20261015_v2.pptx: the 2026-10-15 deck plus the KP-Tracer tumour-level validation.

Four slides go in before "Summary and next steps", each duplicated from the deck's own
title-only slide so that the header, logo and footer band are the deck's:

1. KP-Tracer: Independent Lineage Evidence: schematic of the data and the two information
   streams (expression into ConfidenceOT; CRISPR edits -> inferred tree -> family labels
   as the validation reference). Observed boxes are solid, inferred ones dashed.
2. Tumor-Level Validation Strategy: schematic of the per-pair runs and d_M. The dots are
   illustrative, and the slide says so.
3. ConfidenceOT Recovers Tumor-Family Concordance: native bar chart of the 31 lesion-level
   d_M (sorted), read from metastasis_contrast_source_ranknm256_noscale_ds_cos.csv, and the
   lesion-level numbers from readout_summary.json.
4. Tumor-Family Concordance Across Mice: native chart of every lesion's d_M by mouse with
   the mouse median; the mouse-level numbers and the raw branch beside it.

Explanations go in the notes. One line is added to the summary. The input deck is
only read.
"""
import copy
import json
import subprocess
import sys
from pathlib import Path

import pandas as pd
from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_MARKER_STYLE, XL_TICK_LABEL_POSITION, XL_TICK_MARK
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

SRC = Path(r"D:\Xia_lab\Progress_report\progress_report7\ConfidenceOT_20261015.pptx")
OUT = Path(r"D:\Xia_lab\Progress_report\progress_report7\ConfidenceOT_20261015_v2.pptx")
READOUT = Path(r"D:\Xia_lab\CellOT\results\kp_tracer_lineage_validation\export\readout_tumour_level")
SKILL = Path(r"C:\Users\ncsub\AppData\Local\Packages\Claude_pzs8sxrjxfjjc\LocalCache\Roaming\Claude\local-agent-mode-sessions\skills-plugin\978dbd18-4847-4760-bead-e9e479ea21ba\2ba671bc-afbd-4700-8595-68a5b64c5cef\skills\pptx")
TMP = OUT.with_name("_tmp_v2.pptx")
MAIN, RAW = "ranknm256_noscale_ds_cos", "raw"

# The deck's own colours (theme and existing slides) and the project's gate colours.
ORANGE, LIGHT, FILL = RGBColor(0xBF, 0x57, 0x00), RGBColor(0xEB, 0x8F, 0x2F), RGBColor(0xFD, 0xEB, 0xDD)
INK, MUTED, GREY, PALE = RGBColor(0, 0, 0), RGBColor(0x59, 0x59, 0x59), RGBColor(0xA6, 0xA6, 0xA6), RGBColor(0xF2, 0xF2, 0xF2)
RETAINED, REJECTED = RGBColor(0x2B, 0x8C, 0xBE), RGBColor(0xD7, 0x30, 0x1F)
FONT = "Calibri"


# --- structure ------------------------------------------------------------------------

def duplicate_slides() -> None:
    """Four copies of the title-only slide (slide13.xml, 'Real data rerun'), before the summary."""
    script = SKILL / "scripts" / "add_slide.py"
    subprocess.run([sys.executable, str(script), str(SRC), "slide13.xml", "--after", "slide21.xml", "-o", str(TMP)],
                   check=True, capture_output=True, text=True, encoding="utf-8")
    last = "slide21.xml"
    for _ in range(3):
        names = sorted(Path(n).name for n in __import__("zipfile").ZipFile(TMP).namelist()
                       if n.startswith("ppt/slides/slide") and n.endswith(".xml"))
        newest = max(names, key=lambda n: int(n[5:-4]))
        last = newest
        subprocess.run([sys.executable, str(script), str(TMP), "slide13.xml", "--after", last, "-o", str(TMP)],
                       check=True, capture_output=True, text=True, encoding="utf-8")


def clear_content(slide) -> None:
    keep = {"Rectangle 5", "Rectangle 6", "Picture 8", "TextBox 3"}
    for shape in list(slide.shapes):
        if shape.name not in keep:
            shape._element.getparent().remove(shape._element)


def set_title(slide, text: str) -> None:
    box = next(s for s in slide.shapes if s.name == "TextBox 3")
    runs = [r for p in box.text_frame.paragraphs for r in p.runs]
    runs[0].text = text
    for extra in runs[1:]:
        extra.text = ""


# --- drawing helpers -------------------------------------------------------------------

def text(slide, x, y, w, h, value, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         italic=False, name=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(frame, side, Inches(0.02))
    lines = value if isinstance(value, list) else [value]
    for k, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if k == 0 else frame.add_paragraph()
        paragraph.alignment = align
        parts = line if isinstance(line, list) else [(line, {})]
        for part, style in parts:
            run = paragraph.add_run()
            run.text = part
            font = run.font
            font.name = FONT
            font.size = Pt(style.get("size", size))
            font.bold = style.get("bold", bold)
            font.italic = style.get("italic", italic)
            font.color.rgb = style.get("color", color)
            if "baseline" in style:  # e.g. 30000 superscript, -25000 subscript
                font._element.set("baseline", str(style["baseline"]))
    if name:
        box.name = name
    return box


def box(slide, x, y, w, h, label=None, fill=None, line=INK, dashed=False, size=13, bold=False, color=INK,
        shape=MSO_SHAPE.ROUNDED_RECTANGLE, line_width=1.25, name=None):
    item = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        item.adjustments[0] = 0.12
    if fill is None:
        item.fill.background()
    else:
        item.fill.solid()
        item.fill.fore_color.rgb = fill
    if line is None:
        item.line.fill.background()
    else:
        item.line.color.rgb = line
        item.line.width = Pt(line_width)
        if dashed:
            item.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    item.shadow.inherit = False
    frame = item.text_frame
    frame.word_wrap = True
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(frame, side, Inches(0.05))
    if label is not None:
        lines = label if isinstance(label, list) else [label]
        for k, line_text in enumerate(lines):
            paragraph = frame.paragraphs[0] if k == 0 else frame.add_paragraph()
            paragraph.alignment = PP_ALIGN.CENTER
            run = paragraph.add_run()
            run.text = line_text
            run.font.name, run.font.size, run.font.color.rgb = FONT, Pt(size if k == 0 else size - 2), color
            run.font.bold = bold if k == 0 else False
    if name:
        item.name = name
    return item


def arrow(slide, x1, y1, x2, y2, color=MUTED, width=1.75, dashed=False):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = color
    line.line.width = Pt(width)
    if dashed:
        line.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    tail = etree.SubElement(line.line._get_or_add_ln(), "{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd")
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")
    return line


def segment(slide, x1, y1, x2, y2, color=INK, width=1.25):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = color
    line.line.width = Pt(width)
    return line


def dots(slide, x, y, columns, rows, colours, size=0.13, gap=0.06):
    for k, colour in enumerate(colours):
        r, c = divmod(k, columns)
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + c * (size + gap)), Inches(y + r * (size + gap)),
                                     Inches(size), Inches(size))
        dot.fill.solid()
        dot.fill.fore_color.rgb = colour
        dot.line.fill.background()
        dot.shadow.inherit = False


def notes(slide, value: str) -> None:
    slide.notes_slide.notes_text_frame.text = value


# --- slide 1: the data -----------------------------------------------------------------

def slide_data(slide) -> None:
    # Left: one mouse, several primaries and metastases (schematic).
    box(slide, 0.45, 0.95, 3.75, 3.15, fill=PALE, line=None, name="Mouse panel")
    text(slide, 0.6, 1.02, 3.5, 0.35, "One KP mouse", size=15, bold=True)
    box(slide, 0.65, 1.45, 2.15, 2.45, fill=RGBColor(0xFF, 0xFF, 0xFF), line=GREY, name="Lung")
    text(slide, 0.7, 1.5, 2.0, 0.3, "lung", size=12, color=MUTED)
    for (x, y, label) in ((0.85, 1.95, "T1"), (1.75, 1.85, "T2"), (1.2, 2.85, "T3")):
        box(slide, x, y, 0.75, 0.75, label, fill=FILL, line=ORANGE, shape=MSO_SHAPE.OVAL, size=13, bold=True)
    for (x, y, label) in ((3.05, 1.6, "LN"), (3.05, 2.55, "K"), (3.05, 3.35, "S")):
        box(slide, x, y, 0.62, 0.55, label, fill=RGBColor(0xFF, 0xFF, 0xFF), line=ORANGE, shape=MSO_SHAPE.OVAL,
            size=12, bold=True, dashed=False)
    text(slide, 0.65, 3.92, 3.5, 0.3, "primary tumours        metastases", size=12, color=MUTED)

    # Right: two information streams per cell.
    text(slide, 4.6, 0.95, 8.3, 0.35, "Two measurements on the same single cells", size=15, bold=True)
    y1, y2 = 1.55, 3.05
    box(slide, 4.6, y1, 2.45, 0.85, ["scRNA-seq", "gene expression per cell"], size=14, bold=True, name="RNA")
    arrow(slide, 7.1, y1 + 0.42, 7.75, y1 + 0.42)
    box(slide, 7.8, y1, 5.05, 0.85, ["ConfidenceOT input", "expression only: gates and coupling"], fill=FILL,
        line=ORANGE, size=14, bold=True, name="OT input")

    box(slide, 4.6, y2, 2.45, 1.0, ["CRISPR-Cas9 edits", "target-site indels per cell"], size=14, bold=True,
        name="Edits")
    arrow(slide, 7.1, y2 + 0.5, 7.5, y2 + 0.5)
    tree = box(slide, 7.55, y2, 2.35, 1.0, None, dashed=True, name="Tree")
    text(slide, 7.6, y2 + 0.03, 2.25, 0.3, "lineage tree", size=14, bold=True, align=PP_ALIGN.CENTER)
    text(slide, 7.6, y2 + 0.68, 2.25, 0.3, "inferred (Cassiopeia)", size=12, color=MUTED, align=PP_ALIGN.CENTER)
    # A five-leaf cartoon tree inside the dashed box.
    base, top, xs = y2 + 0.66, y2 + 0.38, [8.05, 8.35, 8.65, 9.05, 9.4]
    segment(slide, 8.2, top, 9.2, top)
    for x in (8.2, 9.2):
        segment(slide, x, top, x, top + 0.08)
    segment(slide, 8.05, top + 0.08, 8.35, top + 0.08)
    segment(slide, 9.05, top + 0.08, 9.4, top + 0.08)
    for x in xs:
        segment(slide, x, top + 0.08, x, base)
    segment(slide, 8.65, top, 8.65, base)
    segment(slide, 8.7, top - 0.1, 8.7, top)
    arrow(slide, 9.95, y2 + 0.5, 10.35, y2 + 0.5)
    box(slide, 10.4, y2, 2.45, 1.0, ["Tumour and family labels", "authors' MetFamily"], dashed=True, size=14,
        bold=True, name="Family labels")
    box(slide, 4.6, 4.45, 8.25, 0.7, "Independent validation reference: compared with ConfidenceOT only after it has run",
        fill=PALE, line=None, size=14, bold=True, name="Reference")
    arrow(slide, 11.6, y2 + 1.02, 11.6, 4.42)

    text(slide, 4.6, 2.5, 8.3, 0.4, [[("Lineage information never enters ConfidenceOT", {"bold": True, "color": REJECTED})]],
         size=13, align=PP_ALIGN.LEFT)

    # Legend: observed vs inferred.
    box(slide, 4.6, 5.55, 0.45, 0.3, None, name="Observed key")
    text(slide, 5.12, 5.53, 2.6, 0.35, "experimentally observed", size=12, color=MUTED)
    box(slide, 7.7, 5.55, 0.45, 0.3, None, dashed=True, name="Inferred key")
    text(slide, 8.22, 5.53, 3.0, 0.35, "computationally inferred", size=12, color=MUTED)
    text(slide, 0.45, 6.55, 12.4, 0.35, "Schematic. Data: KP-Tracer (Yang, Jones et al., Cell 2022; Zenodo 5847462).",
         size=11, color=MUTED, italic=True)
    notes(slide, "\n".join([
        "KP-Tracer (Yang, Jones et al., Cell 2022): KP mouse lung adenocarcinoma with a Cas9 lineage recorder. "
        "Each mouse carries several primary tumours and, in some mice, metastases (LN lymph, K kidney, S soft tissue, "
        "and others).",
        "The same single cells are measured twice: scRNA-seq (104,197 cells with integer counts) and the CRISPR "
        "target sites, whose indels are the observed lineage record.",
        "Trees are inferred from those edits with Cassiopeia (old release) and carry no branch lengths; they are a "
        "reconstruction, not ground truth. The metastatic-family labels (MetFamily) are the authors' annotation.",
        "Here ConfidenceOT sees expression only. Lineage, trees and family labels are used only afterwards, as an "
        "independent reference.",
        "The schematic is not data: the numbers of tumours and metastases differ between mice.",
    ]))


# --- slide 2: the strategy --------------------------------------------------------------

def slide_strategy(slide) -> None:
    rows = (("P1", 0.72, "same family"), ("P2", 0.30, "other family"), ("P3", 0.35, "other family"))
    text(slide, 0.45, 0.9, 7.0, 0.35, "For one metastasis M, every primary tumour of the same mouse",
         size=15, bold=True)
    y0, mx, my, mh = 1.45, 5.55, 2.55, 1.6
    for k, (name, share, family) in enumerate(rows):
        y = y0 + k * 1.55
        box(slide, 0.45, y, 2.85, 1.25, None, line=ORANGE, name=f"{name} box")
        text(slide, 0.55, y + 0.06, 0.6, 0.3, name, size=14, bold=True)
        box(slide, 1.55, y + 0.08, 1.6, 0.32, family, fill=FILL if family == "same family" else None,
            line=ORANGE, dashed=True, size=11, name=f"{name} family")
        n = 24
        retained = round(share * n)
        dots(slide, 0.6, y + 0.52, 8, 3, [RETAINED] * retained + [REJECTED] * (n - retained))
        y_end = my + 0.45 + 0.35 * k
        arrow(slide, 3.35, y + 0.62, mx - 0.05, y_end)
        xm, ym = (3.35 + mx) / 2 - 0.3, (y + 0.62 + y_end) / 2 - 0.2
        box(slide, xm, ym, 0.6, 0.4, f"r{k + 1}", fill=RGBColor(0xFF, 0xFF, 0xFF), line=GREY, size=14, bold=True,
            name=f"r{k + 1}")
    text(slide, 3.4, 1.35, 2.1, 0.55, ["one ConfidenceOT run", "per pair"], size=11, color=MUTED,
         align=PP_ALIGN.CENTER)
    box(slide, mx, my, mh, mh, None, line=ORANGE, name="M box")
    text(slide, mx + 0.1, my + 0.05, 1.4, 0.3, "M", size=14, bold=True)
    dots(slide, mx + 0.2, my + 0.45, 6, 5, [GREY] * 30, size=0.12, gap=0.06)
    text(slide, mx - 0.1, my + mh + 0.05, mh + 0.2, 0.55, ["metastatic", "malignant cells"], size=11, color=MUTED,
         align=PP_ALIGN.CENTER)
    text(slide, 3.4, 5.55, 4.2, 0.35, [[("r", {"bold": True}), (" = retained fraction of the primary’s cells",
                                                                 {"color": MUTED})]], size=12)

    # The contrast.
    box(slide, 7.75, 1.45, 5.1, 4.35, None, fill=FILL, line=None, name="Contrast panel")
    text(slide, 7.95, 1.7, 4.8, 1.0, [[("d", {"bold": True, "size": 26}),
                                       ("M", {"bold": True, "size": 26, "baseline": -25000}),
                                       ("  =  mean r(same family)", {"bold": True, "size": 20})],
                                      [("        −  mean r(other family)", {"bold": True, "size": 20})]], size=20)
    text(slide, 7.95, 3.05, 4.7, 2.6, [[("d", {}), ("M", {"baseline": -25000}),
                                        (" > 0: the same-family primary keeps more of its cells for M", {})],
                                       "",
                                       "Family labels come from the lineage data and are attached after the runs; "
                                       "ConfidenceOT never sees them"], size=14)

    # Legend and caption.
    dots(slide, 0.5, 6.15, 1, 1, [RETAINED])
    text(slide, 0.7, 6.08, 2.2, 0.3, "retained source cell", size=12, color=MUTED)
    dots(slide, 2.9, 6.15, 1, 1, [REJECTED])
    text(slide, 3.1, 6.08, 2.2, 0.3, "rejected source cell", size=12, color=MUTED)
    text(slide, 0.45, 6.55, 12.4, 0.35, "Schematic: the dots and shares are illustrative, not data. "
         "Source: primary malignant cells; target: metastatic malignant cells.", size=11, color=MUTED, italic=True)
    notes(slide, "\n".join([
        "For each metastatic lesion M, every eligible primary tumour of the same mouse is paired with M (at least "
        "20 malignant cells per side; 104 eligible pairs of 131).",
        "Each pair is a separate ConfidenceOT run: source = the primary's malignant cells, target = M's malignant "
        "cells (malignant = not is_normal, the authors' call). Representation ranknm256_noscale_ds_cos, M4-E gate, "
        "within-side acceptance 0.99, no cell cap.",
        "Per pair: the retained fraction r of the primary's malignant cells.",
        "Only after the runs, the authors' MetFamily labels split the primaries of that mouse into same-family and "
        "other-family. d_M = mean r of same-family primaries minus mean r of other-family primaries.",
        "A lesion enters the test when its mouse has at least one eligible primary of each kind: 31 lesions in 8 "
        "mice. Prespecified one-sided Wilcoxon of d_M > 0 across lesions, supported if median d > 0 and p < 0.05 "
        "(READOUT_PRESPECIFICATION_2026-10-14.md, committed before the OT ran).",
        "What it does not show: same-family membership is a tumour-level relation. It does not say that any sampled "
        "primary cell seeded M, or by which route.",
    ]))


# --- slide 3: lesion-level result -------------------------------------------------------

def axis_title(axis, rest: str) -> None:
    axis.has_title = True
    paragraph = axis.axis_title.text_frame.paragraphs[0]
    for content, baseline in (("d", None), ("M", -25000), (" " + rest, None)):
        run = paragraph.add_run()
        run.text = content
        run.font.size, run.font.bold, run.font.color.rgb, run.font.name = Pt(13), False, INK, FONT
        if baseline:
            run.font._element.set("baseline", str(baseline))


def chart_style(chart, size=12):
    chart.font.name = FONT
    chart.font.size = Pt(size)
    chart.has_legend = False


def slide_lesions(slide, lesions: pd.DataFrame, summary: dict) -> None:
    ordered = lesions.sort_values("d", ascending=False).reset_index(drop=True)
    data = CategoryChartData()
    data.categories = list(ordered["metastasis"])
    data.add_series("d_M", [round(float(v), 4) for v in ordered["d"]])
    frame = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.45), Inches(0.95), Inches(8.4),
                                   Inches(5.45), data)
    frame.name = "Lesion d chart"
    chart = frame.chart
    chart_style(chart)
    chart.has_title = False
    plot = chart.plots[0]
    plot.gap_width = 40
    series = plot.series[0]
    series.invert_if_negative = False
    for k, value in enumerate(ordered["d"]):
        point = series.points[k]
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = ORANGE if value > 0 else GREY
        point.format.line.fill.background()
    # PowerPoint also inverts a negative point unless the point itself says otherwise.
    namespace = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
    for dpt in series._element.iter(f"{namespace}dPt"):
        if dpt.find(f"{namespace}invertIfNegative") is None:
            flag = etree.Element(f"{namespace}invertIfNegative")
            flag.set("val", "0")
            dpt.find(f"{namespace}idx").addnext(flag)
    category = chart.category_axis
    category.tick_label_position = XL_TICK_LABEL_POSITION.NONE
    category.major_tick_mark = XL_TICK_MARK.NONE
    category.format.line.color.rgb = INK
    category.format.line.width = Pt(1.25)
    value = chart.value_axis
    value.has_major_gridlines = True
    value.major_gridlines.format.line.color.rgb = RGBColor(0xE6, 0xE6, 0xE6)
    value.format.line.fill.background()
    value.tick_labels.font.size = Pt(12)
    value.tick_labels.font.color.rgb = MUTED
    axis_title(value, "(same − other family)")
    text(slide, 0.6, 6.42, 8.2, 0.3, "31 metastatic lesions, sorted; each bar is one lesion", size=12, color=MUTED,
         align=PP_ALIGN.CENTER)

    # Numbers.
    x = 9.2
    text(slide, x, 1.0, 3.7, 0.9, [[(f"{summary['positive_d']} / {summary['metastases']}", {"bold": True, "size": 40,
                                                                                         "color": ORANGE})]])
    text(slide, x, 1.85, 3.7, 0.6, [[("lesions with d", {}), ("M", {"baseline": -25000}),
                                     (f" > 0  ({summary['negative_d']} < 0, {summary['zero_d']} = 0)", {})]],
         size=13, color=MUTED)
    text(slide, x, 2.55, 3.7, 1.3, [[("median d", {"color": MUTED}), ("M", {"color": MUTED, "baseline": -25000}),
                                     (f"   +{summary['median_d']:.3f}", {"bold": True})],
                                    [("mean d", {"color": MUTED}), ("M", {"color": MUTED, "baseline": -25000}),
                                     (f"   +{summary['mean_d']:.3f}", {"bold": True})]],
         size=18)
    p = summary["wilcoxon_p_one_sided"]
    mantissa, exponent = f"{p:.1e}".split("e")
    text(slide, x, 3.75, 3.7, 1.2, [[("one-sided Wilcoxon, lesion level", {"color": MUTED, "size": 13})],
                                    [(f"p = {mantissa} × 10", {"bold": True}),
                                     (f"−{abs(int(exponent))}", {"bold": True, "baseline": 30000})]],
         size=18)
    text(slide, x, 5.0, 3.7, 1.3, ["104 OT pairs · 31 lesions · 8 mice", "ranknm256_noscale_ds_cos, M4-E",
                                   "lesions are nested in mice (next slide)"], size=12, color=MUTED)
    text(slide, 0.45, 6.75, 12.4, 0.3, "Source: kp_tracer_confidenceot/readout (02_kptracer_readout.py), "
         "metastasis_contrast_source_ranknm256_noscale_ds_cos.csv and readout_summary.json.", size=10,
         color=MUTED, italic=True)
    notes(slide, "\n".join([
        "Result of the prespecified tumour-level readout, main representation ranknm256_noscale_ds_cos.",
        f"104 primary-metastasis OT pairs completed; {summary['metastases']} lesions in {summary['mice']} mice enter "
        f"the test. d_M > 0 in {summary['positive_d']}, < 0 in {summary['negative_d']}, = 0 in {summary['zero_d']}.",
        f"Median d_M +{summary['median_d']:.4f}, mean +{summary['mean_d']:.4f}; one-sided Wilcoxon across lesions "
        f"p = {p:.5e}. Supported by the prespecified criterion.",
        "This p value treats lesions as units. Lesions of one mouse share their primaries and are not independent "
        "animals; the next slide shows the mouse level.",
        "Reading: when ConfidenceOT gates a primary against a metastasis, it keeps more of the primary's cells when "
        "the two tumours belong to the same metastatic family, by the authors' lineage annotation.",
        "It does not show which primary cells seeded the metastasis. The cell-level lineage analysis (E1) is a "
        "separate, preregistered endpoint whose result is inconclusive; Level 2/3 analyses are still running.",
    ]))


# --- slide 4: mouse level ----------------------------------------------------------------

def slide_mice(slide, lesions: pd.DataFrame, summary: dict, raw: dict) -> None:
    mice = sorted(lesions["patient_id"].astype(str).unique(), key=int)
    groups = {m: lesions.loc[lesions["patient_id"].astype(str).eq(m), "d"].tolist() for m in mice}
    depth = max(len(v) for v in groups.values())
    data = CategoryChartData()
    data.categories = mice
    for k in range(depth):
        data.add_series(f"lesion {k + 1}", [round(float(groups[m][k]), 4) if k < len(groups[m]) else None
                                            for m in mice])
    data.add_series("mouse median", [round(float(summary["per_mouse_median_d"][m]), 4) for m in mice])
    frame = slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(0.45), Inches(0.95), Inches(8.4),
                                   Inches(5.45), data)
    frame.name = "Mouse d chart"
    chart = frame.chart
    chart_style(chart)
    for index, series in enumerate(chart.plots[0].series):
        series.format.line.fill.background()
        series.smooth = False
        marker = series.marker
        if index < depth:
            marker.style = XL_MARKER_STYLE.CIRCLE
            marker.size = 8
            marker.format.fill.solid()
            marker.format.fill.fore_color.rgb = LIGHT
            marker.format.line.fill.background()
        else:
            marker.style = XL_MARKER_STYLE.DASH
            marker.size = 20
            marker.format.fill.solid()
            marker.format.fill.fore_color.rgb = INK
            marker.format.line.color.rgb = INK
    category = chart.category_axis
    category.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    category.major_tick_mark = XL_TICK_MARK.NONE
    category.tick_labels.font.size = Pt(13)
    category.format.line.color.rgb = INK
    category.format.line.width = Pt(1.25)
    category.has_title = True
    category.axis_title.text_frame.text = "mouse"
    run = category.axis_title.text_frame.paragraphs[0].runs[0]
    run.font.size, run.font.bold, run.font.color.rgb, run.font.name = Pt(13), False, INK, FONT
    value = chart.value_axis
    value.has_major_gridlines = True
    value.major_gridlines.format.line.color.rgb = RGBColor(0xE6, 0xE6, 0xE6)
    value.format.line.fill.background()
    value.tick_labels.font.size = Pt(12)
    value.tick_labels.font.color.rgb = MUTED
    axis_title(value, "(same − other family)")
    # Legend, drawn so that it matches the markers exactly.
    dots(slide, 1.4, 6.5, 1, 1, [LIGHT], size=0.12)
    text(slide, 1.6, 6.42, 1.6, 0.3, "one lesion", size=12, color=MUTED)
    segment(slide, 3.15, 6.56, 3.45, 6.56, color=INK, width=2.5)
    text(slide, 3.55, 6.42, 2.0, 0.3, "mouse median", size=12, color=MUTED)
    counts = ", ".join(f"{m}: {len(groups[m])}" for m in mice)

    x = 9.2
    text(slide, x, 1.0, 3.7, 0.9, [[(f"{summary['mice_with_median_d_above_0']} / {summary['mice']}",
                                     {"bold": True, "size": 40, "color": ORANGE})]])
    text(slide, x, 1.85, 3.7, 0.6, [[("mice with median d", {}), ("M", {"baseline": -25000}), (" > 0", {})]],
         size=13, color=MUTED)
    text(slide, x, 2.55, 3.7, 1.2, [[("sign test, mouse level", {"color": MUTED, "size": 13})],
                                    [("p = 0.0039", {"bold": True}), ("  exploratory", {"color": MUTED, "size": 13})]],
         size=18)
    box(slide, x, 3.85, 3.65, 1.2, None, fill=PALE, line=None, name="Raw panel")
    text(slide, x + 0.12, 3.92, 3.4, 1.1, [[("raw expression branch", {"bold": True, "size": 13})],
                                            [("median d", {"color": MUTED}), ("M", {"color": MUTED, "baseline": -25000}),
                                             (f"   +{raw['median_d']:.3f}", {"bold": True})],
                                            [(f"{raw['positive_d']} / {raw['metastases']} lesions, "
                                              f"{raw['mice_with_median_d_above_0']} / {raw['mice']} mice > 0",
                                              {"color": MUTED})]], size=14)
    text(slide, x, 5.2, 3.7, 0.6, "the mouse is the independent unit", size=12, color=MUTED)
    text(slide, 0.45, 6.75, 12.4, 0.3, "Source: metastasis_contrast_source_*.csv and readout_summary.json "
         "(kp_tracer_confidenceot/readout).", size=10, color=MUTED, italic=True)
    notes(slide, "\n".join([
        "Mouse-level view of the same result. Lesions are nested in mice: the lesions of one mouse share their "
        "primary tumours, so the 31 lesions are not 31 independent animals.",
        f"Lesions per mouse: {counts}. Mice 3466 and 3513 contribute one lesion each.",
        f"Every one of the {summary['mice']} mice has a positive median d_M. A one-sided sign test on 8 of 8 mice gives "
        "p = 1/256 = 0.0039. This mouse-level test was not prespecified: it is reported as exploratory.",
        f"Raw expression branch (same pairs, same gate, no rank encoding or depth equalisation): median d_M "
        f"+{raw['median_d']:.4f}, {raw['positive_d']} positive, {raw['negative_d']} negative, {raw['zero_d']} zero; "
        f"{raw['mice_with_median_d_above_0']} of {raw['mice']} mice positive. Raw retains about 99% of source cells "
        "in most pairs, so its contrast has little room.",
        "The cell-level analyses (E1 inconclusive; Level 2 and Level 3 running on TACC) are not part of this result.",
    ]))


def add_summary_line(slide) -> None:
    shape = next(s for s in slide.shapes if s.name == "Text summary real data")
    body = shape.text_frame._txBody
    paragraphs = body.findall("{http://schemas.openxmlformats.org/drawingml/2006/main}p")
    new = copy.deepcopy(paragraphs[-1])
    runs = new.findall("{http://schemas.openxmlformats.org/drawingml/2006/main}r")
    for extra in runs[1:]:
        new.remove(extra)
    runs[0].find("{http://schemas.openxmlformats.org/drawingml/2006/main}t").text = (
        "KP-Tracer: same-family primaries retain more source cells (27 of 31 lesions, 8 of 8 mice)")
    body.append(new)


def main():
    lesions_all = pd.read_csv(READOUT / f"metastasis_contrast_source_{MAIN}.csv")
    lesions = lesions_all[lesions_all["enters_test"].astype(bool)].copy()
    summary_all = json.loads((READOUT / "readout_summary.json").read_text())
    summary = summary_all["labels"][MAIN]["primary_source"]
    raw = summary_all["labels"][RAW]["primary_source"]
    assert len(lesions) == summary["metastases"] == 31 and lesions["patient_id"].nunique() == summary["mice"] == 8
    assert abs(lesions["d"].median() - summary["median_d"]) < 1e-12

    duplicate_slides()
    prs = Presentation(str(TMP))
    slides = list(prs.slides)
    new = slides[19:23]  # inserted after the 19th slide, before the summary
    for slide, title in zip(new, ("KP-Tracer: Independent Lineage Evidence", "Tumor-Level Validation Strategy",
                                  "ConfidenceOT Recovers Tumor-Family Concordance",
                                  "Tumor-Family Concordance Across Mice")):
        clear_content(slide)
        set_title(slide, title)
    slide_data(new[0])
    slide_strategy(new[1])
    slide_lesions(new[2], lesions, summary)
    slide_mice(new[3], lesions, summary, raw)
    add_summary_line(slides[23])
    for number, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if shape.name == "Rectangle 6":
                for field in shape._element.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}fld"):
                    field.find("{http://schemas.openxmlformats.org/drawingml/2006/main}t").text = str(number)
    prs.save(str(OUT))
    TMP.unlink()
    print(OUT, len(prs.slides), "slides")


if __name__ == "__main__":
    main()
